#!/usr/bin/env python3
"""MCP Server for file sharing with session-based URLs"""

import os
import sys
import json
import uuid
import shutil
import base64
from datetime import datetime, timedelta
from typing import List, Dict, Any, Optional
from urllib.parse import quote

from mcp.server import Server
from mcp.server.stdio import stdio_server
import mcp.types as types

BASE_DIR = "/files"
SESSIONS_DIR = os.path.join(BASE_DIR, "sessions")
SHARED_DIR = os.path.join(BASE_DIR, "shared")
SESSION_FILE = os.path.join(BASE_DIR, "sessions.json")
DOMAIN = "https://files.netwize.work"
SESSION_TTL_HOURS = 24

os.makedirs(SESSIONS_DIR, exist_ok=True)
os.makedirs(SHARED_DIR, exist_ok=True)

app = Server("file-share")


def make_link(url: str, text: str = "Скачать") -> str:
    return f"[{text}]({url})"


def build_download_url(session_id: str, filename: str) -> str:
    raw = f"{session_id}/{filename}".encode('utf-8')
    key = base64.urlsafe_b64encode(raw).decode('ascii').rstrip('=')
    return f"{DOMAIN}/d/{quote(key, safe='')}"


def load_sessions() -> Dict:
    if os.path.exists(SESSION_FILE):
        with open(SESSION_FILE, 'r') as f:
            return json.load(f)
    return {}


def save_sessions(sessions: Dict):
    with open(SESSION_FILE, 'w') as f:
        json.dump(sessions, f, indent=2)


def cleanup_expired_sessions():
    sessions = load_sessions()
    now = datetime.now()
    expired = []
    for sid, data in sessions.items():
        expires = datetime.fromisoformat(data['expires'])
        if now > expires:
            session_path = os.path.join(SESSIONS_DIR, sid)
            if os.path.exists(session_path):
                shutil.rmtree(session_path)
            expired.append(sid)
    for sid in expired:
        del sessions[sid]
    if expired:
        save_sessions(sessions)
    return len(expired)


@app.list_tools()
async def list_tools() -> List[types.Tool]:
    return [
        types.Tool(
            name="share_file",
            description="Share an existing file and get a download link. Returns URL for downloading.",
            inputSchema={
                "type": "object",
                "properties": {
                    "file_path": {"type": "string", "description": "Path to the file to share"},
                    "session_id": {"type": "string", "description": "Existing session ID (optional, creates new if not provided)"},
                    "description": {"type": "string", "description": "Description for new session (optional)"}
                },
                "required": ["file_path"]
            }
        ),
        types.Tool(
            name="create_and_share",
            description="Create a new file with content and share it. Returns download URL.",
            inputSchema={
                "type": "object",
                "properties": {
                    "filename": {"type": "string", "description": "Name for the new file"},
                    "content": {"type": "string", "description": "Text content for the file"},
                    "session_id": {"type": "string", "description": "Existing session ID (optional)"},
                    "description": {"type": "string", "description": "Description for new session (optional)"}
                },
                "required": ["filename", "content"]
            }
        ),
        types.Tool(
            name="create_binary_and_share",
            description="Create a binary file (docx, xlsx, pdf) from JSON data and share it",
            inputSchema={
                "type": "object",
                "properties": {
                    "filename": {"type": "string", "description": "Name for the file (with extension)"},
                    "file_type": {"type": "string", "enum": ["docx", "xlsx", "pdf", "txt", "csv", "json"], "description": "Type of file to create"},
                    "content": {"type": "string", "description": "Content (text for txt/pdf/docx, JSON for xlsx/csv/json)"},
                    "session_id": {"type": "string", "description": "Existing session ID (optional)"}
                },
                "required": ["filename", "file_type", "content"]
            }
        ),
        types.Tool(
            name="create_session",
            description="Create a new sharing session. Returns session ID.",
            inputSchema={
                "type": "object",
                "properties": {
                    "description": {"type": "string", "description": "Description for the session"}
                },
                "required": []
            }
        ),
        types.Tool(
            name="list_sessions",
            description="List all active sharing sessions",
            inputSchema={
                "type": "object",
                "properties": {},
                "required": []
            }
        ),
        types.Tool(
            name="get_session_links",
            description="Get all download links for a session",
            inputSchema={
                "type": "object",
                "properties": {
                    "session_id": {"type": "string", "description": "Session ID"}
                },
                "required": ["session_id"]
            }
        ),
        types.Tool(
            name="cleanup_sessions",
            description="Remove expired sessions and free up space",
            inputSchema={
                "type": "object",
                "properties": {},
                "required": []
            }
        ),
        types.Tool(
            name="share_directory",
            description="Share all files from a directory",
            inputSchema={
                "type": "object",
                "properties": {
                    "directory_path": {"type": "string", "description": "Path to directory"},
                    "pattern": {"type": "string", "description": "File pattern filter (e.g., '*.pdf')"},
                    "description": {"type": "string", "description": "Session description"}
                },
                "required": ["directory_path"]
            }
        )
    ]


@app.call_tool()
async def call_tool(name: str, arguments: Dict[str, Any]) -> List[types.TextContent]:
    try:
        cleanup_expired_sessions()
        
        if name == "create_session":
            description = arguments.get("description", "")
            sid = str(uuid.uuid4())
            session_path = os.path.join(SESSIONS_DIR, sid)
            os.makedirs(session_path, exist_ok=True)
            sessions = load_sessions()
            sessions[sid] = {
                'created': datetime.now().isoformat(),
                'expires': (datetime.now() + timedelta(hours=SESSION_TTL_HOURS)).isoformat(),
                'description': description,
                'files': []
            }
            save_sessions(sessions)
            return [types.TextContent(type="text", text=f"Session created: {sid}\nExpires: {sessions[sid]['expires'][:19]}")]
        
        elif name == "share_file":
            file_path = arguments["file_path"]
            if not os.path.exists(file_path):
                return [types.TextContent(type="text", text=f"Error: File not found: {file_path}")]
            
            filename = os.path.basename(file_path)
            sid = arguments.get("session_id")
            description = arguments.get("description", f"Shared file: {filename}")
            
            if not sid:
                sid = str(uuid.uuid4())
                session_path = os.path.join(SESSIONS_DIR, sid)
                os.makedirs(session_path, exist_ok=True)
                sessions = load_sessions()
                sessions[sid] = {
                    'created': datetime.now().isoformat(),
                    'expires': (datetime.now() + timedelta(hours=SESSION_TTL_HOURS)).isoformat(),
                    'description': description,
                    'files': []
                }
                save_sessions(sessions)
            
            session_path = os.path.join(SESSIONS_DIR, sid)
            if not os.path.exists(session_path):
                return [types.TextContent(type="text", text=f"Error: Session not found: {sid}")]
            
            dest_path = os.path.join(session_path, filename)
            shutil.copy2(file_path, dest_path)
            
            sessions = load_sessions()
            if sid in sessions and filename not in sessions[sid]['files']:
                sessions[sid]['files'].append(filename)
                save_sessions(sessions)
            
            url = build_download_url(sid, filename)
            return [types.TextContent(type="text", text=make_link(url, "Скачать"))]
        
        elif name == "create_and_share":
            filename = arguments["filename"]
            content = arguments["content"]
            sid = arguments.get("session_id")
            description = arguments.get("description", f"Created file: {filename}")
            
            if not sid:
                sid = str(uuid.uuid4())
                session_path = os.path.join(SESSIONS_DIR, sid)
                os.makedirs(session_path, exist_ok=True)
                sessions = load_sessions()
                sessions[sid] = {
                    'created': datetime.now().isoformat(),
                    'expires': (datetime.now() + timedelta(hours=SESSION_TTL_HOURS)).isoformat(),
                    'description': description,
                    'files': []
                }
                save_sessions(sessions)
            
            session_path = os.path.join(SESSIONS_DIR, sid)
            file_path = os.path.join(session_path, filename)
            
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(content)
            
            sessions = load_sessions()
            if sid in sessions and filename not in sessions[sid]['files']:
                sessions[sid]['files'].append(filename)
                save_sessions(sessions)
            
            url = build_download_url(sid, filename)
            return [types.TextContent(type="text", text=make_link(url, "Скачать"))]
        
        elif name == "create_binary_and_share":
            filename = arguments["filename"]
            file_type = arguments["file_type"]
            content = arguments["content"]
            sid = arguments.get("session_id")
            
            if not sid:
                sid = str(uuid.uuid4())
                session_path = os.path.join(SESSIONS_DIR, sid)
                os.makedirs(session_path, exist_ok=True)
                sessions = load_sessions()
                sessions[sid] = {
                    'created': datetime.now().isoformat(),
                    'expires': (datetime.now() + timedelta(hours=SESSION_TTL_HOURS)).isoformat(),
                    'description': f"Binary file: {filename}",
                    'files': []
                }
                save_sessions(sessions)
            
            session_path = os.path.join(SESSIONS_DIR, sid)
            file_path = os.path.join(session_path, filename)
            
            if file_type == "txt":
                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write(content)
            elif file_type == "json":
                data = json.loads(content)
                with open(file_path, 'w', encoding='utf-8') as f:
                    json.dump(data, f, indent=2)
            elif file_type == "csv":
                import csv
                data = json.loads(content)
                with open(file_path, 'w', newline='', encoding='utf-8') as f:
                    if isinstance(data, list) and data:
                        writer = csv.DictWriter(f, fieldnames=data[0].keys())
                        writer.writeheader()
                        writer.writerows(data)
            elif file_type == "docx":
                from docx import Document
                doc = Document()
                for para in content.split('\n'):
                    if para.strip():
                        doc.add_paragraph(para)
                doc.save(file_path)
            elif file_type == "xlsx":
                from openpyxl import Workbook
                data = json.loads(content)
                wb = Workbook()
                ws = wb.active or wb.create_sheet()
                if isinstance(data, list):
                    for row in data:
                        if isinstance(row, dict):
                            ws.append(list(row.values()))
                        else:
                            ws.append(row)
                wb.save(file_path)
            elif file_type == "pdf":
                from reportlab.lib.pagesizes import letter
                from reportlab.pdfgen import canvas
                c = canvas.Canvas(file_path, pagesize=letter)
                y = 750
                for line in content.split('\n'):
                    c.drawString(50, y, line[:100])
                    y -= 15
                    if y < 50:
                        c.showPage()
                        y = 750
                c.save()
            
            url = build_download_url(sid, filename)
            return [types.TextContent(type="text", text=make_link(url, "Скачать"))]
        
        elif name == "list_sessions":
            sessions = load_sessions()
            if not sessions:
                return [types.TextContent(type="text", text="No active sessions")]
            
            result = "Active Sessions:\n\n"
            for sid, data in sessions.items():
                result += f"Session: {sid}\n"
                result += f"  Created: {data['created'][:19]}\n"
                result += f"  Expires: {data['expires'][:19]}\n"
                result += f"  Files: {len(data['files'])}\n"
                result += f"  Description: {data.get('description', 'N/A')}\n\n"
            
            return [types.TextContent(type="text", text=result)]
        
        elif name == "get_session_links":
            sid = arguments["session_id"]
            sessions = load_sessions()
            if sid not in sessions:
                return [types.TextContent(type="text", text=f"Session not found: {sid}")]
            
            data = sessions[sid]
            result = f"Session: {sid}\n"
            result += f"Expires: {data['expires'][:19]}\n\n"
            result += "Download Links:\n\n"
            
            for fname in data['files']:
                url = build_download_url(sid, fname)
                result += f"- {fname}: {make_link(url, 'Скачать')}\n"
            
            return [types.TextContent(type="text", text=result)]
        
        elif name == "cleanup_sessions":
            count = cleanup_expired_sessions()
            return [types.TextContent(type="text", text=f"Cleaned up {count} expired sessions")]
        
        elif name == "share_directory":
            dir_path = arguments["directory_path"]
            pattern = arguments.get("pattern", "*")
            
            if not os.path.isdir(dir_path):
                return [types.TextContent(type="text", text=f"Error: Directory not found: {dir_path}")]
            
            import glob
            files = glob.glob(os.path.join(dir_path, pattern))
            if not files:
                return [types.TextContent(type="text", text=f"No files matching pattern: {pattern}")]
            
            sid = str(uuid.uuid4())
            session_path = os.path.join(SESSIONS_DIR, sid)
            os.makedirs(session_path, exist_ok=True)
            
            sessions = load_sessions()
            sessions[sid] = {
                'created': datetime.now().isoformat(),
                'expires': (datetime.now() + timedelta(hours=SESSION_TTL_HOURS)).isoformat(),
                'description': arguments.get("description", f"Shared directory: {dir_path}"),
                'files': []
            }
            
            for f in files:
                if os.path.isfile(f):
                    filename = os.path.basename(f)
                    dest = os.path.join(session_path, filename)
                    shutil.copy2(f, dest)
                    sessions[sid]['files'].append(filename)
            
            save_sessions(sessions)
            
            result = f"Shared {len(sessions[sid]['files'])} files!\n\n"
            for fname in sessions[sid]['files']:
                url = build_download_url(sid, fname)
                result += f"- {fname}: {make_link(url, 'Скачать')}\n"
            
            return [types.TextContent(type="text", text=result)]
        
        else:
            return [types.TextContent(type="text", text=f"Unknown tool: {name}")]
    
    except Exception as e:
        return [types.TextContent(type="text", text=f"Error: {str(e)}")]


async def main():
    async with stdio_server() as (read_stream, write_stream):
        await app.run(read_stream, write_stream, app.create_initialization_options())


if __name__ == "__main__":
    import asyncio
    asyncio.run(main())
